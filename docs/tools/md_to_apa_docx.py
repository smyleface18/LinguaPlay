"""Convierte los documentos Markdown de docs/ a Word (.docx) con formato APA 7.

Uso (desde la raíz del repositorio):
    python -m pip install python-docx
    python docs/tools/md_to_apa_docx.py docs/manual-tecnico.md docs/entregables/Manual-tecnico-LinguaPlay.docx

Formato aplicado (APA 7.ª edición, trabajo de estudiante):
- Portada con título, autores, institución, curso, docente y fecha (del
  encabezado YAML del Markdown).
- Tabla de contenido (campo TOC de Word: se actualiza al abrir el documento).
- Times New Roman 12, interlineado doble, márgenes de 2,54 cm, sangría de
  primera línea de 1,27 cm y número de página arriba a la derecha.
- Títulos: nivel 1 centrado en negrita; nivel 2 alineado a la izquierda en
  negrita; nivel 3 a la izquierda en negrita y cursiva.
- Tablas y figuras: "**Tabla N**" / "**Figura N**" en negrita, el título en
  cursiva en la línea siguiente y la nota ("*Nota.* ...") debajo.
- Referencias en página nueva, con sangría francesa de 1,27 cm.

Markdown admitido: encabezado YAML simple (clave: valor), títulos #, ## y ###,
párrafos, listas con viñetas y numeradas (anidadas), tablas con "|", bloques
de código con ```, imágenes ![alt](ruta), **negrita**, *cursiva* y `código`.
Los marcadores [PENDIENTE: ...] y [CAPTURA: ...] se resaltan en amarillo.
"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = 'Times New Roman'
CODE_FONT = 'Courier New'
INDENT = Cm(1.27)
MAX_IMAGE_WIDTH = Cm(16)


# ---------------------------------------------------------------- estilos

def set_font(style_or_run, size=12, name=FONT, bold=None, italic=None):
    font = style_or_run.font
    font.name = name
    font.size = Pt(size)
    font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run.element, 'get_or_add_rPr') else None
    if rpr is not None:
        fonts = rpr.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts')
            rpr.append(fonts)
        for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            fonts.set(qn(attr), name)


def setup_document(doc):
    section = doc.sections[0]
    for side in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
        setattr(section, side, Cm(2.54))

    normal = doc.styles['Normal']
    set_font(normal)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = INDENT

    for level, (align, italic) in {
        1: (WD_ALIGN_PARAGRAPH.CENTER, False),
        2: (WD_ALIGN_PARAGRAPH.LEFT, False),
        3: (WD_ALIGN_PARAGRAPH.LEFT, True),
    }.items():
        style = doc.styles[f'Heading {level}']
        set_font(style, bold=True, italic=italic)
        hpf = style.paragraph_format
        hpf.alignment = align
        hpf.first_line_indent = Cm(0)
        hpf.left_indent = Cm(0)
        hpf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        hpf.space_before = Pt(0)
        hpf.space_after = Pt(0)
        hpf.keep_with_next = True

    # Número de página arriba a la derecha.
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.first_line_indent = Cm(0)
    add_field(header, 'PAGE')

    # Pide a Word actualizar los campos (tabla de contenido) al abrir.
    settings = doc.settings.element
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    settings.append(update)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    set_font(run)
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t')
    text.text = '1' if instruction == 'PAGE' else 'Haga clic derecho y elija "Actualizar campo" para ver la tabla de contenido.'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


# ---------------------------------------------------------------- texto

INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[(?:PENDIENTE|CAPTURA)[^\]]*\])')


def add_inline(paragraph, text, italic=False, size=12):
    text = re.sub(r'\[([^\]]+)\]\((https?://[^)]+)\)', r'\1 (\2)', text)
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=size, bold=True, italic=italic or None)
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            set_font(run, size=size - 1, name=CODE_FONT)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_font(run, size=size, italic=True)
        elif part.startswith('[PENDIENTE') or part.startswith('[CAPTURA'):
            run = paragraph.add_run(part)
            set_font(run, size=size, italic=italic or None)
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            run = paragraph.add_run(part)
            set_font(run, size=size, italic=italic or None)


def paragraph(doc, text, indent=True, align=None, italic=False, keep_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = INDENT if indent else Cm(0)
    if align is not None:
        p.alignment = align
    if keep_next:
        p.paragraph_format.keep_with_next = True
    add_inline(p, text, italic=italic)
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- bloques

def title_page(doc, meta):
    for _ in range(3):
        paragraph(doc, '', indent=False)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    run = title.add_run(meta.get('titulo', 'Documento'))
    set_font(run, bold=True)
    if meta.get('subtitulo'):
        paragraph(doc, meta['subtitulo'], indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(doc, '', indent=False)
    for key in ('autores', 'institucion', 'curso', 'docente', 'fecha', 'version'):
        value = meta.get(key)
        if value:
            label = 'Versión del software: ' if key == 'version' else ''
            paragraph(doc, label + value, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)


def table_of_contents(doc):
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.first_line_indent = Cm(0)
    set_font(heading.add_run('Tabla de contenido'), bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
    page_break(doc)


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn('w:tcBorders'))
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge, value in edges.items():
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), value)
        element.set(qn('w:sz'), '8')
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), '000000')
        borders.append(element)


TEXT_WIDTH_CM = 16.0
MIN_COLUMN_CM = 2.0


def column_widths(rows, columns):
    """Ancho de cada columna proporcional a su texto más largo (con un mínimo)."""
    longest = [
        max(len(re.sub(r'[*`]', '', row[c])) if c < len(row) else 0 for row in rows)
        for c in range(columns)
    ]
    # Que la palabra más larga de la columna quepa sin cortarse (letra de 10 pt).
    longest_word = [
        max(
            (len(word) for row in rows if c < len(row) for word in re.sub(r'[*`]', '', row[c]).split()),
            default=0,
        )
        for c in range(columns)
    ]
    minimums = [max(MIN_COLUMN_CM, 0.21 * n + 0.4) for n in longest_word]
    weights = [min(max(n, 6), 80) for n in longest]
    total = sum(weights)
    widths = [max(TEXT_WIDTH_CM * w / total, m) for w, m in zip(weights, minimums)]
    scale = TEXT_WIDTH_CM / sum(widths)
    return [Cm(w * scale) for w in widths]


def add_table(doc, rows):
    header, *body = rows
    table = doc.add_table(rows=len(rows), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = column_widths(rows, len(header))
    for c, width in enumerate(widths):
        table.columns[c].width = width
    for r, row in enumerate(rows):
        for c in range(len(header)):
            cell = table.cell(r, c)
            cell.width = widths[c]
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after = Pt(2)
            text = row[c] if c < len(row) else ''
            if r == 0:
                set_font(p.add_run(text.replace('**', '')), size=10, bold=True)
            else:
                add_inline(p, text, size=10)
            # Estilo APA: solo líneas horizontales (arriba, bajo el encabezado y al final).
            edges = {'left': 'nil', 'right': 'nil', 'insideV': 'nil'}
            edges['top'] = 'single' if r in (0, 1) else 'nil'
            edges['bottom'] = 'single' if r in (0, len(rows) - 1) else 'nil'
            set_cell_border(cell, **edges)
    return table


def add_image(doc, alt, rel_path, base_dir):
    path = os.path.normpath(os.path.join(base_dir, rel_path))
    if not os.path.exists(path):
        paragraph(doc, f'[PENDIENTE: imagen no encontrada: {rel_path}]', indent=False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    picture = run.add_picture(path)
    if picture.width > MAX_IMAGE_WIDTH:
        ratio = MAX_IMAGE_WIDTH / picture.width
        picture.width = int(picture.width * ratio)
        picture.height = int(picture.height * ratio)


def add_code(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.left_indent = INDENT
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        set_font(p.add_run(line if line else ' '), size=9, name=CODE_FONT)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.first_line_indent = Cm(0)
    spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def add_list_item(doc, marker, text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = INDENT * (level + 1)
    pf.first_line_indent = -Cm(0.63)
    run = p.add_run(f'{marker}\t')
    set_font(run)
    pf.tab_stops.add_tab_stop(INDENT * (level + 1))
    add_inline(p, text)


# ---------------------------------------------------------------- conversión

def list_level(spaces):
    """Nivel de anidamiento según la indentación (2 o 3 espacios por nivel)."""
    if spaces == 0:
        return 0
    return 1 if spaces <= 3 else 2


def parse_front_matter(lines):
    meta = {}
    if lines and lines[0].strip() == '---':
        end = lines.index('---', 1)
        for line in lines[1:end]:
            if ':' in line:
                key, value = line.split(':', 1)
                meta[key.strip()] = value.strip().strip('"')
        lines = lines[end + 1:]
    return meta, lines


def convert(md_path, out_path):
    base_dir = os.path.dirname(os.path.abspath(md_path))
    raw = open(md_path, encoding='utf-8').read().replace('\r\n', '\n').split('\n')
    meta, lines = parse_front_matter(raw)

    doc = Document()
    setup_document(doc)
    title_page(doc, meta)
    table_of_contents(doc)

    i = 0
    buffer = []
    buffer_indent = [0]
    in_references = False

    def flush():
        nonlocal buffer
        if buffer:
            text = ' '.join(s.strip() for s in buffer)
            if in_references:
                p = paragraph(doc, text, indent=False)
                p.paragraph_format.left_indent = INDENT
                p.paragraph_format.first_line_indent = -INDENT
            elif re.fullmatch(r'\*\*(Tabla|Figura) \d+\*\*', text):
                paragraph(doc, text, indent=False, keep_next=True)
            elif re.fullmatch(r'\*[^*]+\*', text) and not text.startswith('*Nota.'):
                paragraph(doc, text, indent=False, keep_next=True)
            elif text.startswith('*Nota.*') or text.startswith('[CAPTURA'):
                paragraph(doc, text, indent=False)
            elif buffer_indent[0] >= 2:
                # Párrafo que continúa un ítem de lista: alineado con su texto.
                p = paragraph(doc, text, indent=False)
                p.paragraph_format.left_indent = INDENT * (list_level(buffer_indent[0]) + 1)
            else:
                paragraph(doc, text)
            buffer = []

    first_h1 = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith('```'):
            flush()
            code = []
            i += 1
            indent = len(line) - len(line.lstrip())
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i][indent:] if lines[i][:indent].strip() == '' else lines[i])
                i += 1
            add_code(doc, code)
            i += 1
            continue

        heading = re.match(r'^(#{1,3})\s+(.*)$', line)
        if heading:
            flush()
            level = len(heading.group(1))
            text = heading.group(2).strip()
            in_references = level == 1 and text.lower() == 'referencias'
            if level == 1 and (in_references and not first_h1):
                page_break(doc)
            first_h1 = False if level == 1 else first_h1
            h = doc.add_heading(level=level)
            add_inline(h, text)
            for run in h.runs:
                run.font.bold = True
                if level == 3:
                    run.font.italic = True
            i += 1
            continue

        if stripped.startswith('|'):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue

        image = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', stripped)
        if image:
            flush()
            add_image(doc, image.group(1), image.group(2), base_dir)
            i += 1
            continue

        item = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', line)
        if item:
            flush()
            level = list_level(len(item.group(1)))
            marker = '•' if item.group(2) in '-*' else item.group(2)
            text = item.group(3)
            i += 1
            # Continuación del ítem en líneas siguientes indentadas (sin ser otro ítem ni código).
            while i < len(lines) and lines[i].strip() and not re.match(r'^\s*([-*]|\d+\.)\s+', lines[i]) \
                    and not lines[i].strip().startswith(('```', '|', '#', '![')) and lines[i].startswith(' '):
                text += ' ' + lines[i].strip()
                i += 1
            add_list_item(doc, marker, text, level)
            continue

        if not stripped:
            flush()
            i += 1
            continue

        if not buffer:
            buffer_indent[0] = len(line) - len(line.lstrip())
        buffer.append(line)
        i += 1

    flush()
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    print(f'generado: {out_path}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
